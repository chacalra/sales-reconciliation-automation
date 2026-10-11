import csv
import shutil
import subprocess
import sys
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from sales_automation import reporting
from sales_automation.loaders import load_source
from sales_automation.models import Counts, Result, Sale
from sales_automation.reconciliation import reconcile
from sales_automation.reporting import OUTPUT_FILES, write_report


@pytest.mark.parametrize(
    "field,value",
    [
        ("sku", "SKU\x01INVALID"),
        ("order_id", "ID\x01INVALID"),
        ("sku", "SKU\x0b"),
        ("sku", "\x1fSKU"),
    ],
)
def test_excel_control_character_is_rejected_without_aborting_cli(config, tmp_path, field, value):
    source = config.sources[0]
    input_path = tmp_path / "sales.csv"
    good = {
        "order_id": "GOOD",
        "sku": "SKU",
        "quantity": "1",
        "unit_price": "10.00",
        "currency": "USD",
        "date": "2026-10-05",
    }
    bad = {**good, "order_id": "BAD", field: value}
    with input_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(good))
        writer.writeheader()
        writer.writerows([good, bad])
    raw = {
        "base_currency": "USD",
        "supported_currencies": ["USD"],
        "output_dir": str(config.output_dir),
        "sources": {
            source.name: {
                "file": str(input_path),
                "date_format": source.date_format,
                "columns": source.columns,
                "amount_type": source.amount_type,
            }
        },
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(raw), encoding="utf-8")
    process = subprocess.run(
        [sys.executable, "-m", "sales_automation", "--config", str(config_path)],
        capture_output=True,
        text=True,
    )
    assert process.returncode == 0, process.stderr
    for filename in (*OUTPUT_FILES, "app.log"):
        assert (config.output_dir / filename).is_file()
    with (config.output_dir / "issues.csv").open(encoding="utf-8-sig") as handle:
        issues = list(csv.DictReader(handle))
    assert len(issues) == 1
    assert issues[0]["category"] == "rejected"
    assert issues[0]["row"] == "3"
    assert f"control character(s) in {field}" in issues[0]["issue"]
    assert repr(value) in issues[0]["issue"]
    if field == "order_id":
        assert issues[0]["order_id"] == "ID\\x01INVALID"
    workbook = load_workbook(config.output_dir / "report.xlsx", data_only=True)
    assert workbook["Summary"]["B5"].value == 2
    assert workbook["Summary"]["B6"].value == 1
    assert workbook["Summary"]["B7"].value == 1
    assert workbook["Transactions"].max_row == 2
    assert workbook["Transactions"]["C2"].value == "GOOD"
    assert workbook["Issues"]["E2"].value == issues[0]["issue"]
    assert "Reconciliation PASS" in (config.output_dir / "app.log").read_text()


@pytest.mark.parametrize("sheet_name", [0, "Sales"])
@pytest.mark.parametrize("duplicate", [False, True])
def test_original_xlsx_headers_checked_before_pandas_renaming(
    config, tmp_path, sheet_name, duplicate
):
    path = tmp_path / "sales.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    headers = ["id", "product_sku", "qty", "total_amount", "currency_code", "created_at"]
    values = ["0001", "SKU", 1, "10.00", "EUR", "2026-10-06T14:30:00"]
    if duplicate:
        headers.append("total_amount")
        values.append("999.00")
    sheet.append(headers)
    sheet.append(values)
    workbook.save(path)
    source = replace(config.sources[1], file=path, sheet_name=sheet_name)
    if duplicate:
        with pytest.raises(ValueError, match="shopify: duplicate column headers"):
            load_source(source)
    else:
        frame = load_source(source)
        assert list(frame.columns) == headers
        assert len(frame) == 1
        assert frame.iloc[0]["id"] == "0001"
        assert frame.iloc[0]["total_amount"] == "10.00"


@pytest.fixture
def report_result():
    sale = Sale(
        "amazon",
        2,
        "OLD",
        "SKU",
        1,
        Decimal("10"),
        "USD",
        date(2026, 10, 5),
        Decimal("10"),
        Decimal("10"),
    )
    return Result(sales=[sale], counts={"amazon": Counts(input_records=1)})


@pytest.mark.parametrize("has_previous", [True, False])
@pytest.mark.parametrize("failed_filename", OUTPUT_FILES)
def test_publication_failure_restores_complete_previous_state(
    tmp_path,
    report_result,
    monkeypatch,
    has_previous,
    failed_filename,
):
    total = reconcile(report_result)
    if has_previous:
        write_report(tmp_path, report_result, total, "USD")
    previous = {path.name: path.read_bytes() for path in tmp_path.iterdir()}
    report_result.sales[0] = replace(report_result.sales[0], order_id="NEW")
    real_replace = Path.replace
    attempted = []

    def fail_publication(path, target):
        if not attempted:
            assert all((path.parent / name).is_file() for name in OUTPUT_FILES)
        attempted.append(path.name)
        if Path(target).name == failed_filename:
            raise OSError("simulated publication failure")
        return real_replace(path, target)

    monkeypatch.setattr(Path, "replace", fail_publication)
    with pytest.raises(OSError, match="previous output state restored"):
        write_report(tmp_path, report_result, total, "USD")
    assert failed_filename in attempted
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == previous


def test_incomplete_staged_outputs_do_not_touch_previous_reports(
    tmp_path, report_result, monkeypatch
):
    total = reconcile(report_result)
    write_report(tmp_path, report_result, total, "USD")
    previous = {path.name: path.read_bytes() for path in tmp_path.iterdir()}
    real_write_csv = reporting.write_csv

    def omit_issues(path, headers, rows):
        if path.name != "issues.csv":
            real_write_csv(path, headers, rows)

    monkeypatch.setattr("sales_automation.reporting.write_csv", omit_issues)
    with pytest.raises(OSError, match="Missing or empty staged output: issues.csv"):
        write_report(tmp_path, report_result, total, "USD")
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == previous


def test_backup_failure_does_not_start_publication(tmp_path, report_result, monkeypatch):
    total = reconcile(report_result)
    write_report(tmp_path, report_result, total, "USD")
    previous = {path.name: path.read_bytes() for path in tmp_path.iterdir()}

    def fail_backup(*args, **kwargs):
        raise OSError("simulated backup failure")

    monkeypatch.setattr("sales_automation.reporting.shutil.copy2", fail_backup)
    with pytest.raises(OSError, match="simulated backup failure"):
        write_report(tmp_path, report_result, total, "USD")
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == previous


def test_failed_rollback_keeps_recovery_backups(tmp_path, report_result, monkeypatch):
    total = reconcile(report_result)
    write_report(tmp_path, report_result, total, "USD")
    previous = {name: (tmp_path / name).read_bytes() for name in OUTPUT_FILES}
    report_result.sales[0] = replace(report_result.sales[0], order_id="NEW")
    real_replace, real_copy = Path.replace, shutil.copy2

    def fail_publish(path, target):
        if Path(target).name == "issues.csv":
            raise OSError("simulated publication failure")
        return real_replace(path, target)

    def fail_restore(source, target):
        if Path(source).parent.name == "backup":
            raise OSError("simulated restore failure")
        return real_copy(source, target)

    monkeypatch.setattr(Path, "replace", fail_publish)
    monkeypatch.setattr("sales_automation.reporting.shutil.copy2", fail_restore)
    with pytest.raises(OSError, match="rollback incomplete.*Recovery backups retained"):
        write_report(tmp_path, report_result, total, "USD")
    (staging,) = tmp_path.glob(".sales-report-*")
    assert {name: (staging / "backup" / name).read_bytes() for name in OUTPUT_FILES} == previous
