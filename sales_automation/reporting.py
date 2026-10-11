"""Readable reports with FX provenance, reconciliation, and formula-safe text."""

import csv
import shutil
from dataclasses import asdict
from decimal import Decimal
from pathlib import Path
from tempfile import mkdtemp

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Counts, Result, Sale
from .reconciliation import revenue
from .validators import escape_excel_controls

ISSUE_COLUMNS = ["source", "row", "order_id", "category", "issue"]
SALE_COLUMNS = list(Sale.__dataclass_fields__) + ["base_currency"]
OUTPUT_FILES = ("report.xlsx", "issues.csv", "consolidated_sales.csv")


def safe_csv(value: object) -> object:
    # Prevent spreadsheet formula injection when a CSV is opened in Excel.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def sale_rows(result: Result, base: str) -> list[list]:
    return [list(asdict(sale).values()) + [base] for sale in result.sales]


def write_csv(path: Path, headers: list[str], rows: list[list]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows([[safe_csv(value) for value in row] for row in rows])


def write_report(output: Path, result: Result, total: Counts, base: str) -> None:
    # Rejected IDs may themselves contain illegal characters. Render them explicitly
    # rather than dropping data or letting the Issues worksheet fail to serialize.
    issue_rows = [
        [
            escape_excel_controls(value) if isinstance(value, str) else value
            for value in asdict(issue).values()
        ]
        for issue in result.issues
    ]
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    summary.append(["SALES RECONCILIATION", base])
    summary.append(["Metric", "Value"])
    metrics = [
        ("Total Orders", len({(sale.source, sale.order_id) for sale in result.sales})),
        (f"Total Revenue ({base})", float(revenue(result))),
        ("Input Records", total.input_records),
        ("Valid Records", total.valid_records),
        ("Rejected Records", total.rejected_records),
        ("Duplicates", total.duplicates),
        ("Reconciliation", "PASS: input = valid + rejected + duplicates"),
        ("Schema Issues", sum(issue.category == "schema" for issue in result.issues)),
        ("FX Origin", result.rates.origin if result.rates else "identity"),
        ("FX Policy", "Latest reference rates; rounded per line; not historical accounting"),
    ]
    for row in metrics:
        summary.append(row)
    summary["B4"].number_format = "#,##0.00"
    summary.append([])
    summary.append(["Revenue by Channel", f"Revenue ({base})", "Valid Lines"])
    chart_start = summary.max_row
    for source in result.counts:
        sales = [sale for sale in result.sales if sale.source == source]
        summary.append(
            [source, float(sum((s.normalized_amount for s in sales), Decimal(0))), len(sales)]
        )
        summary.cell(summary.max_row, 2).number_format = "#,##0.00"
    chart_end = summary.max_row
    chart = BarChart()
    chart.title = f"Revenue by channel ({base})"
    chart.add_data(
        Reference(summary, min_col=2, min_row=chart_start, max_row=chart_end), titles_from_data=True
    )
    chart.set_categories(Reference(summary, min_col=1, min_row=chart_start + 1, max_row=chart_end))
    summary.add_chart(chart, "E2")
    summary.append([])
    summary.append(["Revenue by Currency", "Original Revenue", f"Revenue ({base})"])
    for currency in sorted({sale.currency for sale in result.sales}):
        sales = [sale for sale in result.sales if sale.currency == currency]
        summary.append(
            [
                currency,
                float(sum((s.original_amount for s in sales), Decimal(0))),
                float(sum((s.normalized_amount for s in sales), Decimal(0))),
            ]
        )
        for col in (2, 3):
            summary.cell(summary.max_row, col).number_format = "#,##0.00"
    summary.append([])
    summary.append(["Source Reconciliation", "Input", "Valid", "Rejected", "Duplicates"])
    for source, counts in result.counts.items():
        summary.append([source, *asdict(counts).values()])
    summary.append([])
    summary.append(["FX Currency", f"Units per {base}", "Rate Date", "Origin"])
    if result.rates:
        for code, rate in sorted(result.rates.rates.items()):
            summary.append([code, float(rate), result.rates.dates.get(code), result.rates.origin])
            summary.cell(summary.max_row, 2).number_format = "0.00000000"
    transactions = workbook.create_sheet("Transactions")
    transactions.append(SALE_COLUMNS)
    for row in sale_rows(result, base):
        transactions.append(
            [float(value) if isinstance(value, Decimal) else value for value in row]
        )
    issues = workbook.create_sheet("Issues")
    issues.append(ISSUE_COLUMNS)
    for row in issue_rows:
        issues.append(row)
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        if sheet != summary:
            sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Literal text, even when input starts with '='.
                cell.alignment = Alignment(vertical="top")
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(color="FFFFFF", bold=True)
        for column in sheet.columns:
            width = min(64, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    for row in transactions.iter_rows(min_row=2):
        row[5].number_format = "0.000000"
        row[7].number_format = "yyyy-mm-dd"
        row[8].number_format = "#,##0.00"
        row[9].number_format = "#,##0.00"
    output.mkdir(parents=True, exist_ok=True)
    # One writer per output directory. Roll back ordinary publication failures;
    # this is a logical transaction, not crash-safe filesystem-wide atomicity.
    staging = Path(mkdtemp(prefix=".sales-report-", dir=output))
    preserve_staging = False
    try:
        workbook.save(staging / "report.xlsx")
        write_csv(staging / "issues.csv", ISSUE_COLUMNS, issue_rows)
        write_csv(staging / "consolidated_sales.csv", SALE_COLUMNS, sale_rows(result, base))
        for filename in OUTPUT_FILES:
            path = staging / filename
            if not path.is_file() or path.stat().st_size == 0:
                raise OSError(f"Missing or empty staged output: {filename}")
        backups = staging / "backup"
        backups.mkdir()
        previous = set()
        for filename in OUTPUT_FILES:
            if (output / filename).exists():
                shutil.copy2(output / filename, backups / filename)
                previous.add(filename)
        published = []
        try:
            for filename in OUTPUT_FILES:
                (staging / filename).replace(output / filename)
                published.append(filename)
        except OSError as error:
            rollback_errors = []
            for filename in reversed(published):
                try:
                    if filename in previous:
                        shutil.copy2(backups / filename, output / filename)
                    else:
                        (output / filename).unlink()
                except OSError as rollback_error:
                    rollback_errors.append(f"{filename}: {rollback_error}")
            if rollback_errors:
                preserve_staging = True
                raise OSError(
                    f"Output publication failed: {error}; rollback incomplete: "
                    f"{'; '.join(rollback_errors)}. Recovery backups retained at {backups}"
                ) from error
            raise OSError(
                f"Output publication failed; previous output state restored: {error}"
            ) from error
    finally:
        if not preserve_staging:
            shutil.rmtree(staging)
