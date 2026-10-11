import csv
from dataclasses import replace
from decimal import Decimal
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from sales_automation.duplicate_detector import detect_duplicates
from sales_automation.exchange_rates import parse_response, write_cache
from sales_automation.loaders import load_source
from sales_automation.main import main, run
from sales_automation.models import Counts, Result
from sales_automation.normalizers import normalize_fields
from sales_automation.reconciliation import reconcile, revenue
from sales_automation.validators import validate_record


@pytest.fixture
def valid_fields():
    return {
        "order_id": "A-01",
        "sku": "SKU-01",
        "quantity": "2",
        "amount": "10.50",
        "currency": "USD",
        "sale_date": "2026-10-05",
    }


@pytest.mark.parametrize("index,expected", [(0, 37), (1, 29), (2, 23)])
def test_csv_and_xlsx_loaders(config, index, expected):
    frame = load_source(config.sources[index])
    assert len(frame) == expected
    assert all(isinstance(value, str) for value in frame.iloc[0])


@pytest.mark.parametrize(
    "field,value,message",
    [
        ("quantity", "-3", "invalid quantity"),
        ("quantity", "abc", "invalid quantity"),
        ("quantity", "1.5", "invalid quantity"),
        ("quantity", "0", "invalid quantity"),
        ("amount", "-20", "invalid price"),
        ("amount", "NaN", "invalid price"),
        ("amount", "Infinity", "invalid price"),
        ("currency", "", "missing currency"),
        ("currency", "XYZ", "unknown or unsupported currency"),
        ("order_id", "", "missing order_id"),
        ("sku", "", "missing sku"),
        ("sale_date", "2026-02-30", "invalid sale_date"),
    ],
)
def test_invalid_record(config, valid_fields, field, value, message):
    valid_fields[field] = value
    sale, issue = validate_record(valid_fields, config.sources[0], 2, config.supported_currencies)
    assert sale is None
    assert message in issue.issue


def test_normalize_and_line_total(config):
    source = config.sources[1]
    raw = {
        "id": " 0001 ",
        "product_sku": " SKU-01 ",
        "qty": "3",
        "total_amount": "10.00",
        "currency_code": " eur ",
        "created_at": "2026-10-06T14:30:00",
    }
    fields = normalize_fields(raw, source)
    assert fields["order_id"] == "0001"
    assert fields["currency"] == "EUR"
    sale, issue = validate_record(fields, source, 2, config.supported_currencies)
    assert issue is None
    assert sale.original_amount == Decimal("10.00")  # Never multiply a line total by qty.
    assert abs(sale.unit_price - Decimal(10) / 3) < Decimal("0.000001")


def test_comma_decimal(config):
    raw = {
        "transaction_id": "001",
        "product": "X",
        "units": "2",
        "price": "10,25",
        "currency": "BRL",
        "sale_date": "07/10/2026",
    }
    fields = normalize_fields(raw, config.sources[2])
    sale, issue = validate_record(fields, config.sources[2], 2, config.supported_currencies)
    assert issue is None
    assert sale.original_amount == Decimal("20.50")
    assert sale.sale_date.isoformat() == "2026-10-07"


def test_duplicates_are_scoped_and_conflicts_reported(config, valid_fields):
    sale, _ = validate_record(valid_fields, config.sources[0], 2, config.supported_currencies)
    other_sku = replace(sale, sku="SECOND")
    other_channel = replace(sale, source="shopify")
    conflict = replace(sale, row=5, original_amount=Decimal("999"))
    unique, issues = detect_duplicates(
        [sale, other_sku, other_channel, replace(sale, row=4), conflict]
    )
    assert len(unique) == 3
    assert len(issues) == 2
    assert "conflicting" in issues[1].issue


def seed_cache(config, api_payload):
    write_cache(config.cache_file, parse_response(api_payload, "USD", {"EUR", "GBP", "BRL"}, 14))


def test_full_pipeline_reports_and_known_totals(config, api_payload, monkeypatch):
    # Fixed rates make this a genuine numeric oracle rather than a live API test.
    response = Mock()
    response.json.return_value = api_payload
    monkeypatch.setattr("sales_automation.exchange_rates.requests.get", Mock(return_value=response))
    result = run(config)
    total = reconcile(result)
    assert (total.input_records, total.valid_records, total.rejected_records, total.duplicates) == (
        89,
        76,
        8,
        5,
    )
    # Independent totals: USD 2093.50; EUR 1252.50 => 1565.75 rounded per line;
    # GBP 1283.25 / .5; BRL 1211.00 / 5. Aggregate rounding would differ by 12 cents.
    assert revenue(result) == Decimal("6467.95")
    workbook = load_workbook(config.output_dir / "report.xlsx", data_only=True)
    assert workbook.sheetnames == ["Summary", "Transactions", "Issues"]
    assert workbook["Transactions"].max_row == 77
    assert workbook["Issues"].max_row == 14
    assert workbook["Summary"]["B3"].value == 75  # 76 lines, 75 distinct source/orders.
    assert workbook["Summary"]["B4"].value == 6467.95
    with (config.output_dir / "issues.csv").open(encoding="utf-8-sig") as handle:
        issues = list(csv.DictReader(handle))
    assert len(issues) == 13
    assert {issue["category"] for issue in issues} == {"rejected", "duplicate"}
    with (config.output_dir / "consolidated_sales.csv").open(encoding="utf-8-sig") as handle:
        sales = list(csv.DictReader(handle))
    assert len(sales) == 76
    assert sum(Decimal(s["normalized_amount"]) for s in sales) == revenue(result)


def test_missing_required_column_rejects_source_without_stopping(config, api_payload, tmp_path):
    broken = tmp_path / "broken.csv"
    broken.write_text("sku,quantity,unit_price,currency,date\nX,2,10,USD,2026-10-05\n")
    source = replace(config.sources[0], file=broken)
    config = replace(config, sources=(source, config.sources[1]))
    seed_cache(config, api_payload)
    result = run(config, offline=True)
    assert result.counts["amazon"].rejected_records == 1
    assert result.counts["shopify"].valid_records == 26
    assert result.issues[0].category == "schema"
    assert "order_id" in result.issues[0].issue
    reconcile(result)


def test_reconciliation_detects_loss():
    result = Result(counts={"amazon": Counts(input_records=1)})
    with pytest.raises(ValueError, match="Reconciliation failed"):
        reconcile(result)


def test_missing_file_fails_explicitly(config):
    config = replace(
        config, sources=(replace(config.sources[0], file=config.output_dir / "absent.csv"),)
    )
    with pytest.raises(FileNotFoundError):
        run(config)
    assert not (config.output_dir / "report.xlsx").exists()


def test_formula_input_stays_text(config, valid_fields, tmp_path):
    from sales_automation.reporting import write_report

    valid_fields["order_id"] = '=HYPERLINK("https://example.com")'
    sale, _ = validate_record(valid_fields, config.sources[0], 2, config.supported_currencies)
    result = Result(sales=[sale], counts={"amazon": Counts(input_records=1)})
    total = reconcile(result)
    write_report(tmp_path, result, total, "USD")
    sheet = load_workbook(tmp_path / "report.xlsx")["Transactions"]
    assert sheet["C2"].data_type == "s"
    assert sheet["C2"].value.startswith("=")
    with (tmp_path / "consolidated_sales.csv").open(encoding="utf-8-sig") as handle:
        row = next(csv.DictReader(handle))
    assert row["order_id"].startswith("'=")


def test_cli_failure_is_nonzero(tmp_path):
    assert main(["--config", str(tmp_path / "missing.yaml")]) == 1


def test_malformed_csv_never_silently_discards_data(config, tmp_path):
    malformed = tmp_path / "bad.csv"
    malformed.write_text("a,b\n1,2,3\n")
    with pytest.raises(ValueError, match="malformed CSV"):
        load_source(replace(config.sources[0], file=malformed))
